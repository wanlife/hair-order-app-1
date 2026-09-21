# process_data.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import re
import os

def clean_phone_for_dedup(phone_str):
    if not phone_str:
        return ""
    digits = re.sub(r"\D", "", str(phone_str))
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    return digits

def run_excel_processing(file_path, msg_h_text, msg_k_text, global_seen_phones, global_seen_emails):
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # 1. 解析表头
    headers = [cell.value for cell in ws[1]]
    col_map = {}
    for idx, h in enumerate(headers, start=1):
        if h:
            h_str = str(h).strip()
            col_map[h_str] = idx

    # 定位关键列索引
    phone_col = col_map.get("电话") or col_map.get("Phone") or col_map.get("手机号")
    email_col = col_map.get("邮箱") or col_map.get("Email")
    name_col = col_map.get("姓名") or col_map.get("Name") or col_map.get("客户姓名")

    # 检查并新增 话术列（如果不存在则添加表头）
    h_col_name = "Hello 姓名,"
    k_col_name = "Ambassador 话术"
    
    if h_col_name not in col_map:
        ws.cell(row=1, column=ws.max_column + 1, value=h_col_name)
        col_map[h_col_name] = ws.max_column
    if k_col_name not in col_map:
        ws.cell(row=1, column=ws.max_column + 1, value=k_col_name)
        col_map[k_col_name] = ws.max_column

    col_h_idx = col_map[h_col_name]
    col_k_idx = col_map[k_col_name]

    # 定义重复高亮样式 (黄色/浅红)
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    pink_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # 获取字母列号转义（用于 Excel 动态公式，如 D列）
    name_col_letter = openpyxl.utils.get_column_letter(name_col) if name_col else "D"

    # 2. 逐行生成话术公式与去重高亮
    for row in range(2, ws.max_row + 1):
        # A. 动态写入 H 列与 K 列话术公式
        safe_msg_h = msg_h_text.replace('"', '""')
        safe_msg_k = msg_k_text.replace('"', '""')
        
        ws.cell(row=row, column=col_h_idx).value = f'="Hello "&{name_col_letter}{row}&","&CHAR(10)&CHAR(10)&"{safe_msg_h}"'
        ws.cell(row=row, column=col_k_idx).value = f'="Hello "&{name_col_letter}{row}&","&CHAR(10)&CHAR(10)&"{safe_msg_k}"'

        # B. 电话去重检查与黄色高亮
        if phone_col:
            p_val = ws.cell(row=row, column=phone_col).value
            if p_val:
                p_clean = clean_phone_for_dedup(p_val)
                if p_clean:
                    if p_clean in global_seen_phones:
                        ws.cell(row=row, column=phone_col).fill = yellow_fill
                    else:
                        global_seen_phones.add(p_clean)

        # C. 邮箱去重检查与粉色高亮
        if email_col:
            e_val = ws.cell(row=row, column=email_col).value
            if e_val:
                e_clean = str(e_val).strip().lower()
                if e_clean:
                    if e_clean in global_seen_emails:
                        ws.cell(row=row, column=email_col).fill = pink_fill
                    else:
                        global_seen_emails.add(e_clean)

    # 3. 自动归档保存到输出目录
    out_dir = "处理完成"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, f"已处理+{os.path.basename(file_path)}")
    wb.save(out_file)
    return out_file
